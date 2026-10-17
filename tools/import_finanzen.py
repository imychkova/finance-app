#!/usr/bin/env python3
"""Конвертує тижневі вкладки старої таблиці Finanzen (.xlsx) у CSV для вкладки Transactions.

Нічого не пише в Google: створює data/transactions-import.csv, який імпортується вручну
(Файл → Імпортувати → Додати до поточного аркуша).

Повторюваний імпорт: кожен запис має сталий id (хеш вкладки, дати, назви, суми).
Файл data/imported-ids.txt запам'ятовує вже імпортоване (див. --mark-imported),
тому наступного разу в CSV потраплять лише нові рядки.

Запуск:
  python3 -I tools/import_finanzen.py                    # підсумки + CSV з новими записами
  python3 -I tools/import_finanzen.py --mark-imported    # після успішного імпорту в Sheets
"""
import argparse
import collections
import csv
import datetime
import hashlib
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'data'
LAST_TAB = 'ЮлінО'  # вкладки від початку до цієї (не включно) — тижневі

DATE_RE = re.compile(r'^\s*(\d{1,2})\.(\d{1,2})\.(\d{2})')

# Порядок має значення: перше співпадіння виграє. Шукаємо в «блок + назва» (нижній регістр).
RULES = [
    ('Корм для тварин', ['корм ', 'корм_', 'корм коту', 'корм кеші', 'корм для']),
    ('Пальне', ['бензин', 'заправк', 'дизель', 'окко', 'wog', 'укрнафта', 'socar']),
    ('Алкоголь', ['алкоголь', 'пиво', 'вино', 'вина', 'медовух', 'горілк', 'коньяк', 'віскі', 'самогон']),
    ('Побутова хімія', ['хімія', 'порошок', 'плин для', 'мийн', 'відбілю', 'засіб для', 'фосфат']),
    ('Авто', ['міцубісі', 'паркінг', 'парковк', 'таксі', 'шини', 'колес', 'мийка', 'сто ', 'техогляд', 'страхув']),
    ('Кава', ['кава', 'кави ', 'каву', 'кав\'ярн', 'латте', 'капучіно', 'еспресо', 'coffee']),
    ('Аптека', ['аптека', 'ліки', 'вітамін']),
    ('Підписки й інтернет', ['нетфлікс', 'netflix', 'sweet.tv', 'sweet tv', 'світ тв', 'інтернет', 'youtube', 'spotify', 'підписк']),
    ('Ресторани', ['суші', 'піцц', 'піца', 'pizza', 'кебаб', 'хінкальн', 'обід', 'ресторан', 'кафе', 'мегого',
                   'mcdonald', 'макдональдс', 'ролл', 'бургер', 'пекарн', 'доставк', 'glovo', 'bolt food']),
    ('Дім і сад', ['епіцентр', 'саджанц', 'розсад', 'насінн', 'сад + город', 'сад ', 'город', 'банки скляні',
                   'замок', 'швабра', 'сушарка', 'тирса', 'дрова', 'будова', 'інструмент']),
    ('Подарунки', ['подарунк', 'квіти', 'букет']),
    ('Їжа', ['продукт', 'атб', 'сільпо', 'правдоруб', 'близенько', 'рукавичка', 'арсен', 'ашан', 'аврора',
             'м\'ясо', 'мясо', 'яйця', 'овоч', 'фрукт', 'торт', 'тістечк', 'пляцк', 'мамс', 'сухарик', 'чорниц', 'горох', 'квас', 'м\'яс', 'молоко', 'сир', 'масло', 'олія', 'хліб', 'булочк',
             'булка', 'риба', 'качка', 'качки', 'курк', 'куряч', 'свин', 'індич', 'кров\'янка', 'сарделька',
             'фарш', 'горішк', 'сухофрукт', 'ананас', 'банан', 'полуниц', 'черешн', 'абрикос', 'нектарин',
             'лохин', 'кавун', 'помідор', 'огірк', 'картопл', 'цибул', 'капуст', 'слива', 'яблук', 'груш',
             'спеції', 'квасол', 'вода', 'айран', 'кефір', 'творог', 'лаваш', 'пластівці', 'кукурудз',
             'твікс', 'білка', 'пляцок', 'млинц', 'сирник', 'ролліні', 'багет', 'пакет', 'євро', 'з європи',
             'магазин', 'маркет', 'їжа', 'їжу']),
]


def parse_date(text, last_year=None):
    m = DATE_RE.match(text)
    if not m:
        return None
    d, mo, y = map(int, m.groups())
    year = 2000 + y
    if year < 2024 and last_year:  # опечатка на кшталт 29.10.10
        year = last_year
    try:
        return datetime.date(year, mo, d)
    except ValueError:
        return None


def to_amount(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(' ', '').replace(',', '.'))
        except ValueError:
            return None
    return None


def categorize(block, label):
    hay = (block + ' ' + label).lower() + ' '
    for cat, words in RULES:
        if any(w in hay for w in words):
            return cat
    return 'Інше'


def iso_week(d):
    y, w, _ = d.isocalendar()
    return f'{y}-W{w:02d}'


def read_items(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    tabs = names[:names.index(LAST_TAB)]
    items = []
    last_year = None
    for tab in tabs:
        rows = list(wb[tab].iter_rows(values_only=True))
        if not rows:
            continue
        for k, head in enumerate(rows[0]):
            if not isinstance(head, str) or not head.strip():
                continue
            start = parse_date(head, last_year)
            if start:  # основний тижневий блок: назва в колонці k, примітка k+1, сума k+2
                cur = start
                last_year = cur.year
                for r in rows[1:]:
                    label = r[k] if k < len(r) else None
                    qty = r[k + 1] if k + 1 < len(r) else None
                    amt = to_amount(r[k + 2]) if k + 2 < len(r) else None
                    if not (isinstance(label, str) and label.strip()):
                        continue
                    d2 = parse_date(label, last_year)
                    if d2:
                        cur = d2
                        last_year = cur.year
                        continue
                    if amt is None:
                        continue
                    note = label.strip() + (f' ({str(qty).strip()})' if isinstance(qty, str) and qty.strip() else '')
                    items.append(dict(tab=tab, date=cur, label=label.strip(), note=note, amount=amt, block=''))
            else:  # бічний блок: дата k, примітка k+1, сума k+2
                block = head.strip()
                for r in rows[1:]:
                    d = r[k] if k < len(r) else None
                    note = r[k + 1] if k + 1 < len(r) else None
                    amt = to_amount(r[k + 2]) if k + 2 < len(r) else None
                    if isinstance(d, datetime.datetime) and amt is not None:
                        text = note.strip() if isinstance(note, str) else ''
                        items.append(dict(tab=tab, date=d.date(), label=text, note=(block + ' ' + text).strip(),
                                          amount=amt, block=block))
    return items


def build_rows(items):
    seen = collections.Counter()
    out = []
    for it in items:
        key = f"{it['tab']}|{it['date']}|{it['block']}|{it['label']}|{it['amount']}"
        seen[key] += 1  # однакові рядки в один день — різні записи
        uid = 'imp-' + hashlib.sha1(f'{key}|{seen[key]}'.encode()).hexdigest()[:16]
        cash = 'готівк' in (it['label'] + it['note']).lower()
        note = ' '.join(it['note'].split()) + (' [повернення?]' if it['amount'] < 0 else '')
        if note[:1] in '=+-@':
            note = "'" + note
        d = it['date']
        out.append([uid, d.isoformat(), f"{it['amount']:.2f}", categorize(it['block'], it['label']),
                    'Готівка' if cash else 'Невідомо', 'Разом', note[:200], '', '', '', ''])  # I:K заповнить recomputeKeys()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=str(DATA / 'Finanzen.xlsx'))
    ap.add_argument('--out', default=str(DATA / 'transactions-import.csv'))
    ap.add_argument('--all', action='store_true', help='ігнорувати imported-ids.txt')
    ap.add_argument('--mark-imported', action='store_true', help='запам\'ятати id з останнього CSV як імпортовані')
    args = ap.parse_args()

    state = DATA / 'imported-ids.txt'
    done = set(state.read_text().split()) if state.exists() else set()

    if args.mark_imported:
        with open(args.out, newline='', encoding='utf-8') as f:
            ids = {row[0] for row in csv.reader(f)}
        state.write_text('\n'.join(sorted(done | ids)) + '\n')
        print(f'Позначено як імпортовані: {len(ids)} (усього {len(done | ids)}).')
        return

    rows = build_rows(read_items(args.xlsx))
    new = rows if args.all else [r for r in rows if r[0] not in done]
    with open(args.out, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        # без рядка заголовків: у Transactions він уже є, імпорт «Додати до поточного аркуша»
        w.writerows(new)

    print(f'Знайдено записів: {len(rows)}, у CSV (нові): {len(new)} → {args.out}')
    by_month = collections.defaultdict(float)
    by_cat = collections.defaultdict(float)
    for r in new:
        by_month[r[9]] += float(r[2])
        by_cat[r[3]] += float(r[2])
    print('\nПо місяцях:')
    for m in sorted(by_month):
        print(f'  {m}  {by_month[m]:>10,.0f}')
    print('\nПо категоріях:')
    for c, s in sorted(by_cat.items(), key=lambda x: -x[1]):
        print(f'  {c:<22}{s:>10,.0f}  ({sum(1 for r in new if r[3] == c)} шт)')
    other = collections.Counter(r[6].lower() for r in new if r[3] == 'Інше')
    print(f'\nНевпізнані («Інше»), найчастіші з {sum(other.values())}:')
    for text, n in other.most_common(40):
        print(f'  {n:>3}  {text}')


if __name__ == '__main__':
    sys.exit(main())
